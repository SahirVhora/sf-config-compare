"""Tests for deterministic, explainable migration-readiness findings."""

from copy import deepcopy

import openpyxl

from core.quality_analyzer import analyze_migration_readiness
from core.reporter import generate_excel_report, generate_html_report


def _result_with_diffs():
    return {
        "summary": {
            "entities_only_in_a": 1,
            "entities_only_in_b": 0,
            "entities_in_both": 1,
            "fields_matched": 0,
            "fields_with_diff": 2,
            "fields_only_in_a": 0,
            "fields_only_in_b": 0,
            "picklists_only_in_a": 0,
            "picklists_only_in_b": 0,
            "picklists_shared": 1,
            "missing_values_in_a": 0,
            "missing_values_in_b": 1,
            "value_diffs": 1,
        },
        "entity_diffs": [
            {
                "entity_name": "EmpJob",
                "diff_type": "only_in_a",
                "details": {"entity_label": "Job Information"},
            }
        ],
        "field_diffs": [
            {
                "entity_name": "EmpEmployment",
                "field_id": "startDate",
                "field_label": "Start Date",
                "diff_type": "attribute_mismatch",
                "attribute": "required",
                "value_a": "true",
                "value_b": "false",
            },
            {
                "entity_name": "EmpEmployment",
                "field_id": "assignmentIdExternal",
                "field_label": "Assignment ID",
                "diff_type": "attribute_mismatch",
                "attribute": "max_length",
                "value_a": "100",
                "value_b": "32",
            },
        ],
        "picklist_result": {
            "missing_picklists": [],
            "missing_values": [
                {
                    "picklist_id": "employment-type",
                    "external_code": "CONTINGENT",
                    "label": "Contingent Worker",
                    "status": "ACTIVE",
                    "diff_type": "only_in_a",
                }
            ],
            "value_diffs": [
                {
                    "picklist_id": "employment-type",
                    "external_code": "EMPLOYEE",
                    "label_a": "Employee",
                    "label_b": "Permanent Employee",
                    "field_diffs": [
                        {
                            "field": "label_en",
                            "value_a": "Employee",
                            "value_b": "Permanent Employee",
                        }
                    ],
                }
            ],
        },
    }


def test_analysis_emits_explainable_findings_with_fixed_severity_scale():
    analysis = analyze_migration_readiness("SOURCE", "TARGET", _result_with_diffs())

    assert analysis["analysis_type"] == "deterministic_rules"
    assert analysis["uses_ai"] is False
    assert analysis["summary"] == {
        "total": 5,
        "highest_severity": "high",
        "by_severity": {
            "critical": 0,
            "high": 4,
            "medium": 0,
            "low": 1,
            "info": 0,
        },
    }
    required = {
        "rule_id",
        "severity",
        "object_type",
        "object_id",
        "affected_field",
        "source_value",
        "target_value",
        "reason",
        "recommended_correction",
        "detection_method",
    }
    assert all(required <= finding.keys() for finding in analysis["findings"])
    assert {finding["severity"] for finding in analysis["findings"]} <= {
        "critical",
        "high",
        "medium",
        "low",
        "info",
    }
    assert all(
        finding["detection_method"] == "deterministic_rule"
        for finding in analysis["findings"]
    )
    severity_rank = {
        severity: rank
        for rank, severity in enumerate(("critical", "high", "medium", "low", "info"))
    }
    assert [severity_rank[f["severity"]] for f in analysis["findings"]] == sorted(
        severity_rank[f["severity"]] for f in analysis["findings"]
    )


def test_analysis_explains_source_to_target_required_and_length_risks():
    findings = analyze_migration_readiness("SOURCE", "TARGET", _result_with_diffs())[
        "findings"
    ]

    required = next(f for f in findings if f["affected_field"] == "required")
    assert required["severity"] == "high"
    assert required["object_id"] == "EmpEmployment.startDate"
    assert required["source_value"] == "true"
    assert required["target_value"] == "false"
    assert "SOURCE" in required["reason"]
    assert "TARGET" in required["recommended_correction"]

    length = next(f for f in findings if f["affected_field"] == "max_length")
    assert length["severity"] == "high"
    assert "truncate" in length["reason"].lower()


def test_analysis_is_empty_and_explicit_when_comparison_has_no_diffs():
    result = _result_with_diffs()
    result["entity_diffs"] = []
    result["field_diffs"] = []
    result["picklist_result"] = {
        "missing_picklists": [],
        "missing_values": [],
        "value_diffs": [],
    }

    analysis = analyze_migration_readiness("SOURCE", "TARGET", result)

    assert analysis["findings"] == []
    assert analysis["summary"]["total"] == 0
    assert analysis["summary"]["highest_severity"] is None
    assert analysis["summary"]["by_severity"] == {
        "critical": 0,
        "high": 0,
        "medium": 0,
        "low": 0,
        "info": 0,
    }


def test_reports_surface_actionable_validation_findings(tmp_path, monkeypatch):
    import core.reporter as reporter

    result = _result_with_diffs()
    result["quality_analysis"] = analyze_migration_readiness("SOURCE", "TARGET", result)
    monkeypatch.setattr(reporter, "REPORTS_DIR", tmp_path)

    html = generate_html_report("SOURCE", "TARGET", result)
    path = generate_excel_report(
        "SOURCE",
        "TARGET",
        deepcopy(result),
        {"base_url": "https://source.example.com"},
        {"base_url": "https://target.example.com"},
    )

    assert "Migration Readiness Findings" in html
    assert "Recommended correction" in html
    assert "EmpEmployment.startDate" in html

    workbook = openpyxl.load_workbook(path)
    assert "Validation Findings" in workbook.sheetnames
    summary_rows = list(workbook["Summary"].iter_rows(values_only=True))
    assert ("High severity findings", 4) in [row[:2] for row in summary_rows]
    rows = list(workbook["Validation Findings"].iter_rows(values_only=True))
    assert rows[0][:5] == (
        "Severity",
        "Rule ID",
        "Object Type",
        "Object ID",
        "Affected Field",
    )
    assert any("EmpEmployment.startDate" in row for row in rows[1:])


def test_validation_findings_excel_neutralizes_formula_values(tmp_path, monkeypatch):
    import core.reporter as reporter

    result = _result_with_diffs()
    result["field_diffs"][0]["value_a"] = '=HYPERLINK("https://example.invalid")'
    result["quality_analysis"] = analyze_migration_readiness("SOURCE", "TARGET", result)
    monkeypatch.setattr(reporter, "REPORTS_DIR", tmp_path)

    path = generate_excel_report(
        "SOURCE",
        "TARGET",
        result,
        {"base_url": "https://source.example.com"},
        {"base_url": "https://target.example.com"},
    )

    workbook = openpyxl.load_workbook(path, data_only=False)
    rows = list(workbook["Validation Findings"].iter_rows(values_only=True))
    required_row = next(row for row in rows if row[4] == "required")
    assert required_row[5] == '\'=HYPERLINK("https://example.invalid")'
    field_rows = list(workbook["Field Attr Diffs"].iter_rows(values_only=True))
    required_field_row = next(row for row in field_rows if row[3] == "required")
    assert required_field_row[4] == '\'=HYPERLINK("https://example.invalid")'


def test_validation_findings_html_escapes_tenant_controlled_text():
    result = _result_with_diffs()
    result["entity_diffs"][0]["entity_name"] = "<script>alert(1)</script>"
    result["quality_analysis"] = analyze_migration_readiness("SOURCE", "TARGET", result)

    html = generate_html_report("SOURCE", "TARGET", result)

    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html


def test_compare_api_includes_quality_analysis(client, sample_instances):
    source_id, target_id = sample_instances

    response = client.post(
        "/api/v1/compare",
        json={"instance_a_id": source_id, "instance_b_id": target_id},
    )

    assert response.status_code == 200
    analysis = response.get_json()["quality_analysis"]
    assert analysis["analysis_type"] == "deterministic_rules"
    assert analysis["summary"]["total"] > 0
    assert analysis["findings"][0]["recommended_correction"]
